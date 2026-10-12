from collections import Counter, defaultdict
from datetime import date, datetime
from io import BytesIO

from flask import (
    Blueprint, abort, flash, redirect, render_template, request, send_file, url_for,
)
from flask_login import current_user, login_required

from .. import db
from ..constants import (
    STATE_CHECKED, STATE_DRAFT, STATE_REVISION, STATE_SENT, TEAMS,
)
from ..helpers import format_date, month_end, parse_date, parse_month, spv_required, wa_text
from ..models import Report, User


bp = Blueprint("spv", __name__)


def _team_order(users):
    order = {t: i for i, t in enumerate(TEAMS)}
    return sorted(users, key=lambda u: (order.get(u.team, 99), u.name.lower()))


@bp.route("/spv")
@spv_required
def dashboard():
    d = parse_date(request.args.get("tanggal"), date.today())
    team = request.args.get("tim", "")

    users = User.query.filter_by(active=True).filter(User.team != "Supervisor")
    if team:
        users = users.filter_by(team=team)
    users = _team_order(users.all())

    reports = {r.user_id: r for r in Report.query.filter(
        Report.date == d, Report.user_id.in_([u.id for u in users]))}

    # Dikelompokkan menurut tindakan Spv: perlu dicek dulu, lalu lainnya.
    groups = {"check": [], "revision": [], "checked": [], "missing": []}
    for u in users:
        r = reports.get(u.id)
        if not r or r.state == STATE_DRAFT:
            groups["missing"].append((u, None))
        elif r.state == STATE_SENT:
            groups["check"].append((u, r))
        elif r.state == STATE_REVISION:
            groups["revision"].append((u, r))
        else:
            groups["checked"].append((u, r))
    sent = len(users) - len(groups["missing"])

    # Laporan hari lain yang masih menunggu dicek.
    pending = (Report.query.join(User, Report.user_id == User.id)
               .filter(Report.state == STATE_SENT, Report.date != d, User.active.is_(True),
                       Report.source == "app")
               .order_by(Report.date.desc()).limit(30).all())

    return render_template("spv_dashboard.html", day=d, groups=groups, team=team,
                           total=len(users), sent=sent, pending=pending, today=date.today())


def _same_day(report):
    """Laporan terkirim di tanggal yang sama, urut tim lalu nama."""
    rows = (Report.query.join(User, Report.user_id == User.id)
            .filter(Report.date == report.date, Report.state != STATE_DRAFT).all())
    return sorted(rows, key=lambda r: (TEAMS.index(r.user.team) if r.user.team in TEAMS else 99,
                                       r.user.name.lower()))


def _next_to_check(report):
    """Laporan berikutnya (tanggal sama) yang masih perlu dicek, atau None."""
    rows = _same_day(report)
    ids = [r.id for r in rows]
    pos = ids.index(report.id) if report.id in ids else -1
    for r in rows[pos + 1:] + rows[:max(pos, 0)]:
        if r.id != report.id and r.state == STATE_SENT:
            return r
    return None


@bp.route("/spv/report/<int:report_id>")
@spv_required
def detail(report_id):
    report = db.session.get(Report, report_id) or abort(404)
    if report.state == STATE_DRAFT:
        abort(404)

    rows = _same_day(report)
    ids = [r.id for r in rows]
    pos = ids.index(report.id) if report.id in ids else -1
    prev_id = ids[pos - 1] if pos > 0 else None
    next_id = ids[pos + 1] if 0 <= pos < len(ids) - 1 else None
    remaining = sum(1 for r in rows if r.state == STATE_SENT and r.id != report.id)

    return render_template("spv_detail.html", report=report, prev_id=prev_id, next_id=next_id,
                           position=pos + 1, total=len(rows), remaining=remaining,
                           wa=wa_text(report.user.name, report.date, report.day_status, report.items))


@bp.route("/spv/report/<int:report_id>/review", methods=["POST"])
@spv_required
def review(report_id):
    report = db.session.get(Report, report_id) or abort(404)
    action = request.form.get("action")
    note = request.form.get("note", "").strip()[:2000]

    name = report.user.name
    if action == "check":
        report.state = STATE_CHECKED
        message = (f"Laporan {name} sudah dicek.", "success")
    elif action == "revision":
        if not note:
            flash("Tulis dulu apa yang perlu diperbaiki.", "warning")
            return redirect(url_for("spv.detail", report_id=report.id))
        report.state = STATE_REVISION
        message = (f"Laporan {name} dikembalikan untuk diperbaiki.", "warning")
    elif action == "reopen" and report.state == STATE_CHECKED:
        report.state = STATE_SENT
        message = ("Tanda sudah dicek dibatalkan.", "info")
    else:
        abort(400)

    report.spv_note = note or None
    report.reviewed_at = datetime.now()
    report.reviewed_by_id = current_user.id
    report.edited_after_submit = False
    db.session.commit()

    # Setelah dicek / dikembalikan, langsung ke laporan berikutnya yang
    # masih perlu dicek; kalau sudah habis, kembali ke daftar.
    if action in ("check", "revision"):
        nxt = _next_to_check(report)
        if nxt:
            flash(*message)
            return redirect(url_for("spv.detail", report_id=nxt.id))
        flash(message[0] + " Semua laporan di tanggal ini sudah diperiksa.", "success")
        return redirect(url_for("spv.dashboard", tanggal=report.date.isoformat()))
    flash(*message)
    return redirect(url_for("spv.detail", report_id=report.id))


# ------------------------------------------------------------ rekap

def _rekap_scope():
    today = date.today()
    first = parse_month(request.args.get("dari", ""), today.replace(day=1))
    last_first = parse_month(request.args.get("sampai", ""), today.replace(day=1))
    if last_first < first:
        first, last_first = last_first, first
    end = month_end(last_first)

    query = (Report.query.join(User, Report.user_id == User.id)
             .filter(Report.date >= first, Report.date <= end))

    if current_user.is_spv:
        team = request.args.get("tim", "")
        user_id = request.args.get("staf", type=int)
        if team:
            query = query.filter(User.team == team)
        if user_id:
            query = query.filter(Report.user_id == user_id)
        # Draft belum dikirim tidak ikut dihitung di rekap Spv.
        query = query.filter(Report.state != STATE_DRAFT)
    else:
        team, user_id = "", current_user.id
        query = query.filter(Report.user_id == current_user.id)

    return first, end, team, user_id, query.order_by(Report.date, User.name).all()


@bp.route("/rekap")
@login_required
def rekap():
    first, end, team, user_id, reports = _rekap_scope()

    per_user = defaultdict(lambda: {"days": 0, "items": 0, "rutin": 0, "off": 0, "cats": Counter()})
    cats = Counter()
    cats_rutin = Counter()
    for r in reports:
        row = per_user[r.user]
        if r.day_status:
            row["off"] += 1
        if r.items:
            row["days"] += 1
        for it in r.items:
            row["items"] += 1
            row["cats"][it.get("cat", "Lainnya")] += 1
            cats[it.get("cat", "Lainnya")] += 1
            if it.get("rutin"):
                row["rutin"] += 1
                cats_rutin[it.get("cat", "Lainnya")] += 1

    users = _team_order(list(per_user))
    total = sum(cats.values())
    kpi = {
        "total": total,
        "days": sum(v["days"] for v in per_user.values()),
        "rutin": sum(cats_rutin.values()),
        "people": len(per_user),
    }
    staff_options = _team_order(User.query.filter_by(active=True).all()) if current_user.is_spv else []

    return render_template(
        "rekap.html", first=first, end=end, team=team, user_id=user_id,
        users=users, per_user=per_user, cats=cats.most_common(), cats_rutin=cats_rutin,
        kpi=kpi, staff_options=staff_options,
        dari=first.strftime("%Y-%m"), sampai=end.strftime("%Y-%m"),
    )


@bp.route("/rekap/export.xlsx")
@login_required
def rekap_export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    first, end, _team, _user_id, reports = _rekap_scope()

    wb = Workbook()
    ws = wb.active
    ws.title = "Semua Pekerjaan"
    head = ["Tanggal", "Hari", "Nama", "Tim", "No", "Kategori", "Rutin", "Pekerjaan",
            "Status Hari", "Status Report", "Catatan Spv"]
    ws.append(head)
    for r in reports:
        items = r.items or [{"text": "", "cat": "", "rutin": False}]
        for i, it in enumerate(items, 1):
            ws.append([r.date, format_date(r.date).split(",")[0], r.user.name, r.user.team,
                       i if it.get("text") else None, it.get("cat", ""),
                       "Rutin" if it.get("rutin") else ("Insidental" if it.get("text") else ""),
                       it.get("text", ""), r.day_status, r.state_label, r.spv_note or ""])

    fill = PatternFill("solid", fgColor="EA580C")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    for col, width in zip("ABCDEFGHIJK", [12, 9, 22, 15, 5, 30, 11, 90, 14, 15, 40]):
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD-MM-YYYY"
        row[7].alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("Ringkasan")
    ws2.append(["Nama", "Tim", "Hari kerja", "Pekerjaan", "Rutin", "Insidental", "Hari cuti/sakit/izin"])
    summary = defaultdict(lambda: [0, 0, 0, 0])
    for r in reports:
        s = summary[r.user]
        s[0] += 1 if r.items else 0
        s[1] += len(r.items)
        s[2] += sum(1 for it in r.items if it.get("rutin"))
        s[3] += 1 if r.day_status else 0
    for u in _team_order(list(summary)):
        s = summary[u]
        ws2.append([u.name, u.team, s[0], s[1], s[2], s[1] - s[2], s[3]])
    for c in ws2[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = fill
    ws2.column_dimensions["A"].width = 26
    ws2.column_dimensions["B"].width = 16

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    name = f"Rekap Daily Report HRGA {first:%Y-%m}" + (f" sd {end:%Y-%m}" if end.month != first.month or end.year != first.year else "")
    return send_file(buf, as_attachment=True, download_name=name + ".xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
